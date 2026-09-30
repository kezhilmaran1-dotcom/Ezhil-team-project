import os
import re
from io import BytesIO
from fpdf import FPDF
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'", "\u2019": "'",
        "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-",
        "\u2022": "-", "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return re.sub(r"[^\x09\x0A\x0D\x20-\x7E]", "", text)

def _split_sections(text):
    lines = [line.strip() for line in sanitize_text(text).splitlines()]
    return [line for line in lines if line]

def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")

def format_docx(text: str, doc_type: str, terms: str = "", logo_path: str = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if logo_path and os.path.exists(logo_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_path, width=Inches(1.2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in _split_sections(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        if (
            len(line) < 90
            and (line.isupper() or re.match(r"^(\d+[\.\)]|[A-Z][A-Z\s&-]{3,})$", line))
        ):
            r.bold = True

    if terms.strip():
        doc.add_paragraph()
        heading = doc.add_paragraph()
        rr = heading.add_run("TERMS SUMMARY")
        rr.bold = True
        rr.font.name = "Times New Roman"
        rr.font.size = Pt(12)

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for i, term in enumerate([x.strip() for x in terms.split(";") if x.strip()], 1):
            cells = table.add_row().cells
            cells[0].text = str(i)
            cells[1].text = term

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("LegalEase - AI-Powered Legal Document Generator")
    fr.font.name = "Times New Roman"
    fr.font.size = Pt(9)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()

class LegalEasePDF(FPDF):
    def __init__(self, logo_path=None):
        super().__init__()
        self.logo_path = logo_path

    def header(self):
        if self.logo_path and os.path.exists(self.logo_path):
            try:
                self.image(self.logo_path, x=90, y=8, w=30)
                self.ln(18)
            except Exception:
                self.ln(5)
        else:
            self.ln(5)
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 6, "LegalEase", align="C")
        self.ln(8)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(0, 8, "LegalEase - AI-Powered Legal Document Generator", align="C")

def format_pdf(text: str, doc_type: str, logo_path: str = None) -> bytes:
    pdf = LegalEasePDF(logo_path)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, sanitize_text(doc_type).upper(), align="C")
    pdf.ln(5)

    for line in _split_sections(text):
        if (
            len(line) < 90
            and (line.isupper() or re.match(r"^(\d+[\.\)]|[A-Z][A-Z\s&-]{3,})$", line))
        ):
            pdf.set_font("Helvetica", "B", 11)
        else:
            pdf.set_font("Helvetica", size=10)
        pdf.multi_cell(0, 6, line)
        pdf.ln(1)

    return bytes(pdf.output())
